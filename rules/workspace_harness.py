"""Optional smolagents JSON tool runner; SDK calls stay in detect.client."""

import hashlib
import json
import os
import re
import uuid
from dataclasses import asdict, is_dataclass
from decimal import Decimal, InvalidOperation
from typing import Literal

from pydantic import Field

from contracts.models import Contract
from rules.discovery_models import SearchResult
from rules.legal_sources import fetch_source, validate_url
from rules.workspace_chat import ANSWER_PROMPT, AnswerDraft


class ToolAction(Contract):
    tool: Literal[
        "search_documents",
        "read_document",
        "read_excel",
        "summarize_excel",
        "web_search",
        "final_answer",
    ]
    query: str | None = Field(max_length=500)
    document_id: str | None = Field(max_length=100)
    page: int = Field(ge=1, le=250)
    cell_range: str | None = Field(max_length=30)
    final: AnswerDraft | None


class WorkspaceTools:
    def __init__(self, index, *, search=None):
        self.index, self.search_provider = index, search
        self.retrieved, self.trace = {}, []
        self.web_calls = 0
        self.search_events = []
        self.search_counts_known = True

    def _remember(self, tool, chunks, **metadata):
        for chunk in chunks:
            self.retrieved[chunk["chunk_id"]] = chunk
        self.trace.append({"tool": tool, "hits": len(chunks), **metadata})
        return {"sources": chunks}

    def search_documents(self, query):
        if not 1 <= len(query.strip()) <= 500:
            raise ValueError("Invalid search query.")
        return self._remember("search_documents", self.index.search(query), query=query)

    def read_document(self, document_id, page):
        return self._remember(
            "read_document",
            self.index.read_document(document_id, page),
            document_id=document_id,
            page=page,
            provider="reviewed_ingested_text",
        )

    def read_excel(self, document_id, page, cell_range):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        source = self.index.sources.get(document_id)
        if not source or source.get("suffix") != ".xlsx":
            raise FileNotFoundError("Selected XLSX document not found in this workspace.")
        if not re.fullmatch(r"[A-Z]{1,3}\d{1,7}:[A-Z]{1,3}\d{1,7}", cell_range or ""):
            raise ValueError("Use an explicit Excel range such as A1:D20.")
        first_col, first_row, last_col, last_row = range_boundaries(cell_range)
        if min(first_col, first_row) < 1 or last_col < first_col or last_row < first_row:
            raise ValueError("Invalid Excel range.")
        if (last_col - first_col + 1) * (last_row - first_row + 1) > 100:
            raise ValueError("Excel range exceeds 100-cell budget.")
        if not 1 <= page <= len(source["model_pages"]):
            raise ValueError("Unknown worksheet page.")
        path = (
            self.index.store.folder(self.index.workspace_id)
            / "documents"
            / document_id
            / "original.xlsx"
        )
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            worksheet_path = source["worksheets"][page - 1]["path"]
            sheet = next(
                (sheet for sheet in workbook.worksheets if sheet._worksheet_path == worksheet_path),
                None,
            )
            if sheet is None:
                raise ValueError("Worksheet does not match the reviewed extraction.")
            reviewed = dict(
                line.split("\t", 1)
                for line in source["model_pages"][page - 1].splitlines()
                if "\t" in line
            )
            lines = []
            for row in sheet.iter_rows(
                min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col
            ):
                for cell in row:
                    text = reviewed.get(cell.coordinate, "")
                    if cell.value is None or not text or "[" in text:
                        continue
                    matches = str(cell.value) == text
                    if not matches:
                        try:
                            matches = Decimal(str(cell.value)) == Decimal(text)
                        except InvalidOperation:
                            pass
                    if matches:
                        lines.append(f"{cell.coordinate}\t{text}")
        finally:
            workbook.close()
        chunk = {
            "chunk_id": "CHUNK-"
            + hashlib.sha256(
                f"{document_id}:{source['version']}:{page}:{cell_range}".encode()
            ).hexdigest()[:24],
            "document_id": document_id,
            "title": source["title"],
            "collection": source["collection"],
            "sha256": source["sha256"],
            "version": source["version"],
            "page": page,
            "worksheet": f"Sheet {page}",
            "text": "\n".join(lines),
        }
        return self._remember(
            "read_excel",
            [chunk] if lines else [],
            document_id=document_id,
            page=page,
            cell_range=cell_range,
            provider="openpyxl",
            warning="Only reviewed matching cells; formulas are not recalculated.",
        )

    def summarize_excel(self, document_id):
        """Cover every worksheet using only the already reviewed cell projection."""
        source = self.index.sources.get(document_id)
        if not source or source.get("suffix") != ".xlsx":
            raise FileNotFoundError("Selected XLSX document not found in this workspace.")
        pages = source["model_pages"]
        character_budget = min(1800, 18000 // max(1, len(pages)))
        chunks, coverage, citation_options = [], [], []
        for page, text in enumerate(pages, 1):
            original_lines = set(source["pages"][page - 1].splitlines())
            cells = [
                line
                for line in text.splitlines()
                if re.fullmatch(r"[A-Z]+\d+\t[^\n]+", line)
                and "[" not in line
                and line in original_lines
            ]
            labels, numbers = [], []
            for line in cells:
                (
                    numbers
                    if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?%?", line.split("\t", 1)[1])
                    else labels
                ).append(line)
            # Spread numeric examples over the whole sheet, not just its first index chunks.
            numeric_samples = [
                numbers[min(len(numbers) - 1, i * len(numbers) // min(8, len(numbers)))]
                for i in range(min(8, len(numbers)))
            ]
            chosen = set(labels[:8] + numeric_samples)
            sampled, size = [], 0
            for line in cells:
                if line in chosen and size + len(line) + 1 <= character_budget:
                    sampled.append(line)
                    size += len(line) + 1
            coverage.append(
                {
                    "worksheet": f"Sheet {page}",
                    "page": page,
                    "reviewed_cells_available": len(cells),
                    "sampled_cells": len(sampled),
                    "numeric_cells_available": len(numbers),
                    "coverage": "sampled_reviewed_cells",
                }
            )
            if sampled:
                chunks.append(
                    {
                        "chunk_id": "CHUNK-"
                        + hashlib.sha256(
                            f"overview:{document_id}:{source['sha256']}:{source['version']}:{page}".encode()
                        ).hexdigest()[:24],
                        "document_id": document_id,
                        "title": source["title"],
                        "collection": source["collection"],
                        "sha256": source["sha256"],
                        "version": source["version"],
                        "page": page,
                        "worksheet": f"Sheet {page}",
                        "text": "\n".join(sampled),
                    }
                )
                run_lines = []
                for line in source["pages"][page - 1].splitlines() + [""]:
                    if line in sampled:
                        run_lines.append(line)
                    elif run_lines:
                        citation_options.append(
                            {"chunk_id": chunks[-1]["chunk_id"], "quote": "\n".join(run_lines)}
                        )
                        run_lines = []
        return {
            **self._remember(
                "summarize_excel",
                chunks,
                document_id=document_id,
                worksheets_covered=len(pages),
                provider="reviewed_workbook_projection",
            ),
            "worksheet_count": len(pages),
            "worksheets": coverage,
            "coverage": "sampled_reviewed_cells",
            "citation_options": citation_options,
            "note": "Content summary only; no legal evidence is required. All worksheets were inspected through the reviewed projection; samples are partial, private/unknown text is withheld, and no totals across mixed units were inferred. Copy citation_options quotes verbatim; these may contain multiple contiguous original cell lines. For a multi-sheet summary use at most one citation per worksheet plus one header citation, within ten total. Use read_excel for a specific range if needed.",
        }

    def web_search(self, query):
        # Legal topics legitimately mention tenants, addresses and personal data.
        # Reject concrete identifiers rather than censoring those topic words.
        if re.search(
            r"[\w.+-]+@[\w.-]+|\d(?:[ .()-]*\d){7,}|(?:họ\s*(?:và\s*)?tên|cccd|cmnd|địa\s*chỉ|tài\s*khoản)\s*[:=]\s*\S+",
            query,
            re.IGNORECASE,
        ):
            raise ValueError("Do not send private identifiers in public search queries.")
        if self.web_calls >= 1 or not 1 <= len(query.strip()) <= 500:
            raise ValueError("Public search budget exhausted or invalid query.")
        self.web_calls += 1
        if self.search_provider is None:
            from detect.client import DiscoveryTransport

            self.search_provider = DiscoveryTransport(
                cache_dir=self.index.store.folder(self.index.workspace_id) / "public-search-cache"
            )
        previous_events = len(getattr(self.search_provider, "events", []))
        result = self.search_provider.complete(
            "search",
            {
                "topics": [{"id": "chat", "question": query}],
                "workspace_id": self.index.workspace_id,
            },
            SearchResult.model_json_schema(),
            timeout_s=25,
            web_search=True,
        )
        if hasattr(self.search_provider, "events"):
            self.search_events.extend(self.search_provider.events[previous_events:])
        else:
            self.search_counts_known = False
        sources = []
        for lead in SearchResult.model_validate(result).sources:
            try:
                url = validate_url(lead.url)
            except ValueError:
                continue
            if url not in {item["url"] for item in sources}:
                sources.append({"url": url, "verification_status": "unverified_search_lead"})
            if len(sources) == 2:
                break
        chunks = []
        for lead in sources:
            archive = (
                self.index.store.folder(self.index.workspace_id)
                / "public-search-cache"
                / "originals"
            )
            try:
                source = fetch_source(
                    lead["url"], timeout_s=12, max_bytes=2_000_000, archive_dir=archive
                )
                raw_path = next(
                    (
                        path
                        for suffix in (".pdf", ".html")
                        if (path := archive / (source.sha256 + suffix)).is_file()
                    ),
                    None,
                )
                if raw_path is None:
                    raise ValueError("Official original was not archived.")
                raw = raw_path.read_bytes()
                if hashlib.sha256(raw).hexdigest() != source.sha256:
                    raise ValueError("Archived official source hash mismatch.")
            except Exception:  # noqa: BLE001 - unavailable untrusted sources remain unread leads
                lead["retrieval_status"] = "could not confirm updates from this source"
                continue
            # Only this bounded public-source ingestion may update our source snapshot.
            # External selection/extraction changes must still reject the answer.
            with self.index.store._lock:
                from rules.workspace_index import WorkspaceIndex

                fresh = WorkspaceIndex(self.index.store, self.index.workspace_id)
                if fresh.fingerprint != self.index.fingerprint:
                    raise ValueError("Selected sources changed during public search; ask again.")
                document = self.index.store.add_document(
                    self.index.workspace_id,
                    "legal",
                    "official-" + source.sha256[:16] + raw_path.suffix,
                    raw,
                    origin="discovered",
                    official_url=source.url,
                    extracted_pages=source.pages,
                )
                fresh = WorkspaceIndex(self.index.store, self.index.workspace_id)
                self.index.__dict__.update(fresh.__dict__)
            lead.update(
                document_id=document["id"],
                sha256=document["sha256"],
                retrieval_status="ingested_unverified",
                verification_status="to_be_verified",
            )
            hits = [
                item for item in self.index.search(query) if item["document_id"] == document["id"]
            ]
            chunks.extend(hits[:3] or self.index.read_document(document["id"], 1)[:1])
        self.trace.append(
            {
                "tool": "web_search",
                "query": query,
                "hits": len(sources),
                "provider": "detect.client.DiscoveryTransport",
                "sources": sources,
                "model_events": self.search_events,
            }
        )
        return {
            **self._remember("read_public_sources", chunks),
            "search_leads": sources,
            "note": "Fetched official text is unverified legal evidence; currency and applicability need review. Unread leads cannot support a legal answer.",
        }


def run(index, common, client, *, search=None, max_steps=4):
    """Run the actual OSS harness with an adapter to the existing model boundary."""
    try:
        from smolagents import ChatMessage, Model, ToolCallingAgent, tool
        from smolagents.models import ChatMessageToolCall, ChatMessageToolCallFunction
    except ImportError as exc:
        raise RuntimeError(
            "Install optional rules/tooling-requirements.txt; harness unavailable."
        ) from exc

    if not 1 <= max_steps <= 6:
        raise ValueError("Tool harness allows one to six steps.")
    tools = WorkspaceTools(index, search=search)

    @tool
    def search_documents(query: str) -> str:
        """Search selected ingested workspace text.

        Args:
            query: A concise search phrase.
        """
        return json.dumps(tools.search_documents(query), ensure_ascii=False)

    @tool
    def read_document(document_id: str, page: int) -> str:
        """Read reviewed ingested PDF, Word, text or spreadsheet page content.

        Args:
            document_id: An inventory ID selected for this workspace.
            page: One-based extraction page/sheet number.
        """
        return json.dumps(tools.read_document(document_id, page), ensure_ascii=False)

    @tool
    def read_excel(document_id: str, page: int, cell_range: str) -> str:
        """Read actual XLSX cells through openpyxl, limited to reviewed redacted values.

        Args:
            document_id: A selected XLSX inventory ID.
            page: One-based worksheet page from inventory.
            cell_range: Explicit bounded range such as A1:D20, at most100 cells.
        """
        return json.dumps(tools.read_excel(document_id, page, cell_range), ensure_ascii=False)

    @tool
    def summarize_excel(document_id: str) -> str:
        """Get a bounded content overview of ALL reviewed worksheets in a selected workbook.

        Args:
            document_id: A selected XLSX inventory ID already in the workspace.
        """
        return json.dumps(tools.summarize_excel(document_id), ensure_ascii=False)

    @tool
    def web_search(query: str) -> str:
        """Search current official Vietnamese legal websites through the configured provider.

        Args:
            query: Public legal topic only, without private company/person identifiers.
        """
        return json.dumps(tools.web_search(query), ensure_ascii=False)

    @tool
    def final_answer(answer: str) -> str:
        """Return a structured answer only after checking its exact retrieved quotations.

        Args:
            answer: JSON matching the structured AnswerDraft schema.
        """
        draft = AnswerDraft.model_validate_json(answer)
        for reference in draft.citations:
            chunk = tools.retrieved.get(reference.chunk_id)
            if chunk is None:
                raise ValueError(
                    "Final quote is not retrieved; use an observed citation_options entry."
                )
            try:
                index.citation(chunk, reference.quote)
            except ValueError as exc:
                raise ValueError(
                    "Final quote is not exact. Copy an observed citation_options quote verbatim, including coordinate and tabs; do not add characters."
                ) from exc
        return answer

    class BoundaryModel(Model):
        def __init__(self):
            super().__init__(model_id=getattr(client, "model", None))
            self.outputs, self.steps, self.exhausted = [], 0, False

        def generate(self, messages, **kwargs):
            from rules.workspace_index import WorkspaceIndex

            if WorkspaceIndex(index.store, index.workspace_id).fingerprint != index.fingerprint:
                raise ValueError("Selected sources changed before a model call; ask again.")
            if self.steps >= max_steps:
                self.exhausted = True
                raise RuntimeError("Tool harness model-call budget exhausted.")
            self.steps += 1
            serialized = [
                asdict(message) if is_dataclass(message) else message for message in messages
            ]
            available = [item.name for item in kwargs.get("tools_to_call_from", [])]
            if self.steps == max_steps:
                available = ["final_answer"]
            output = client.generate(
                purpose="workspace_tool_harness_v1",
                payload={
                    **common,
                    "index_sha256": index.fingerprint,
                    "inventory": index.inventory(),
                    "messages": serialized,
                    "tools": available,
                    "steps_remaining": max_steps - self.steps,
                },
                prompt=ANSWER_PROMPT
                + " Choose exactly one allowed tool per step. Inspect actual observations before answering. For an Excel content summary or overview, call summarize_excel ONCE to cover all worksheets, then answer from its reviewed cell observations. This is a content-reading task: relevant business cells are sufficient evidence, no legal clauses or legal verification are required. Do not web_search for a content-only request. Do not refuse the supported partial summary because some fields are redacted or not every cell was read; state the coverage limits and use status answered when supported content is available. Do not let earlier legal questions override the user's current summary request. Return final_answer promptly after sufficient observations. When steps_remaining is zero only final_answer is allowed, with either supported content or a precise missing-content explanation. Copy citation_options quotations verbatim when provided; a quote may contain multiple contiguous original cell lines. For multi-sheet summaries use at most one citation per worksheet plus one header citation, never more than ten; summarize fewer facts if needed. Search leads alone cannot support a legal answer. No Python/code execution is available.",
                schema=ToolAction,
                mode="live",
            )
            if output.status not in ("live", "cached") or output.parsed is None:
                raise RuntimeError("Harness model unavailable; no sample fallback.")
            self.outputs.append(output)
            action = ToolAction.model_validate(output.parsed)
            if action.tool not in available:
                if self.steps == max_steps:
                    self.exhausted = True
                raise ValueError("Model requested an unavailable tool.")
            if action.tool != "final_answer":
                tools.trace.append({"tool": action.tool, "status": "requested"})
            arguments = {}
            if action.tool in ("search_documents", "web_search"):
                arguments = {"query": action.query}
            elif action.tool in ("read_document", "read_excel"):
                arguments = {"document_id": action.document_id, "page": action.page}
                if action.tool == "read_excel":
                    arguments["cell_range"] = action.cell_range
            elif action.tool == "summarize_excel":
                arguments = {"document_id": action.document_id}
            else:
                if action.final is None:
                    raise ValueError("Final answer is missing structured evidence.")
                arguments = {"answer": action.final.model_dump_json()}
            return ChatMessage(
                role="assistant",
                content=None,
                tool_calls=[
                    ChatMessageToolCall(
                        function=ChatMessageToolCallFunction(arguments=arguments, name=action.tool),
                        id=uuid.uuid4().hex,
                        type="function",
                    )
                ],
            )

    model = BoundaryModel()
    allowed = [search_documents, read_document, read_excel, summarize_excel, final_answer]
    if os.getenv("LEXCHECK_WORKSPACE_WEB_SEARCH_ENABLED", "0") == "1":
        allowed.append(web_search)
    agent = ToolCallingAgent(
        model=model, tools=allowed, max_steps=max_steps, add_base_tools=False, verbosity_level=-1
    )
    try:
        result = agent.run(
            json.dumps({**common, "inventory": index.inventory()}, ensure_ascii=False)
        )
    except Exception as exc:
        if model.exhausted:
            raise RuntimeError(
                "Tool harness budget exhausted without a grounded final answer."
            ) from exc
        raise RuntimeError("Tool harness unavailable; no sample answer substituted.") from exc
    if model.exhausted:
        raise RuntimeError("Tool harness budget exhausted without a grounded final answer.")
    draft = AnswerDraft.model_validate_json(str(result))
    agent_calls = sum(item.calls for item in model.outputs)
    search_calls = (
        sum(not event.get("cached", False) for event in tools.search_events)
        if tools.search_counts_known
        else None
    )
    provenance = {
        "model": getattr(client, "model", None),
        "status": "live" if any(item.status == "live" for item in model.outputs) else "cached",
        "model_calls": agent_calls + search_calls if search_calls is not None else None,
        "agent_model_calls": agent_calls,
        "public_search_model_calls": search_calls,
        "public_search_events": tools.search_events,
        "index_sha256": index.fingerprint,
        "cache_keys": [item.key for item in model.outputs],
        "retrieval": "model_selected_oss_tools",
        "harness": "smolagents.ToolCallingAgent",
        "harness_version": "1.26.0",
        "steps": model.steps,
        "public_search_requests": tools.web_calls,
    }
    return draft, list(tools.retrieved.values()), tools.trace, provenance
